"""`documents`: produce Markdown, CSV, JSON, Mermaid, SVG, DOCX, PDF, XLSX and PPTX files from structured
content, write them under `documents/` in the workspace and upload them as artifacts (docs/04 §1).

DOCX uses python-docx and PDF uses reportlab when they are installed; XLSX uses openpyxl when installed and
otherwise a small built-in writer; PPTX needs python-pptx. A format whose library is missing returns a clear
error instead of a broken file.
"""

from __future__ import annotations

import csv
import importlib.util
import io
import json
import re
import zipfile
from typing import Any
from xml.sax.saxutils import escape

from z_common.models.common import FieldError
from z_common.models.enums import ArtifactKind

from z_runner.artifacts import ArtifactRefused
from z_runner.client import ApiError
from z_runner.tools.base import ArgumentError, ToolContext, ToolResult, ZTool
from z_runner.workspace import JailError

EXTENSIONS = {
    "md": "md",
    "csv": "csv",
    "json": "json",
    "mermaid": "mmd",
    "svg": "svg",
    "docx": "docx",
    "pdf": "pdf",
    "xlsx": "xlsx",
    "pptx": "pptx",
}
_KIND = {"csv": ArtifactKind.data, "json": ArtifactKind.data, "svg": ArtifactKind.image}
_CONTENT_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
}
_SVG_FORBIDDEN = re.compile(
    r"(?i)<\s*script|\bon[a-z]+\s*=|javascript:|<\s*foreignObject|xlink:href\s*=\s*['\"]\s*(?:https?:|data:)"
)

Scalar = str | int | float | bool | None


def library_available(module: str) -> bool:
    return importlib.util.find_spec(module) is not None


def _cell(value: Scalar) -> str:
    return "" if value is None else str(value)


def _sections(args: dict[str, Any]) -> list[dict[str, Any]]:
    return list(args.get("sections") or [])


def render_markdown(title: str | None, sections: list[dict[str, Any]]) -> str:
    out: list[str] = []
    if title:
        out += [f"# {title}", ""]
    for section in sections:
        if section.get("heading"):
            out += [f"## {section['heading']}", ""]
        for paragraph in section.get("paragraphs") or []:
            out += [paragraph, ""]
        for bullet in section.get("bullets") or []:
            out.append(f"- {bullet}")
        if section.get("bullets"):
            out.append("")
        table = section.get("table")
        if table:
            headers = [_cell(h) for h in table.get("headers") or []]
            if headers:
                out.append("| " + " | ".join(headers) + " |")
                out.append("|" + "---|" * len(headers))
            for row in table.get("rows") or []:
                out.append("| " + " | ".join(_cell(c).replace("|", "\\|") for c in row) + " |")
            out.append("")
    return "\n".join(out).rstrip() + "\n"


def render_csv(headers: list[str], rows: list[list[Scalar]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    if headers:
        writer.writerow(headers)
    writer.writerows([[_cell(c) for c in row] for row in rows])
    return buf.getvalue().encode("utf-8")


def render_docx(title: str | None, sections: list[dict[str, Any]]) -> bytes:
    import docx  # type: ignore[import-untyped,unused-ignore]

    document = docx.Document()
    if title:
        document.add_heading(title, 0)
    for section in sections:
        if section.get("heading"):
            document.add_heading(section["heading"], 1)
        for paragraph in section.get("paragraphs") or []:
            document.add_paragraph(paragraph)
        for bullet in section.get("bullets") or []:
            document.add_paragraph(bullet, style="List Bullet")
        table = section.get("table")
        if table:
            headers = [_cell(h) for h in table.get("headers") or []]
            rows = [[_cell(c) for c in row] for row in table.get("rows") or []]
            cols = max([len(headers)] + [len(r) for r in rows]) if (headers or rows) else 0
            if cols:
                grid = document.add_table(rows=0, cols=cols)
                grid.style = "Table Grid"
                for values in ([headers] if headers else []) + rows:
                    cells = grid.add_row().cells
                    for i, value in enumerate(values[:cols]):
                        cells[i].text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def render_pdf(title: str | None, sections: list[dict[str, Any]]) -> bytes:
    from reportlab.lib import colors  # type: ignore[import-untyped,unused-ignore]
    from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped,unused-ignore]
    from reportlab.lib.styles import getSampleStyleSheet  # type: ignore[import-untyped,unused-ignore]
    from reportlab.platypus import (  # type: ignore[import-untyped,unused-ignore]
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    story: list[Any] = []
    if title:
        story += [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
    for section in sections:
        if section.get("heading"):
            story.append(Paragraph(escape(section["heading"]), styles["Heading2"]))
        for paragraph in section.get("paragraphs") or []:
            story += [Paragraph(escape(paragraph), styles["BodyText"]), Spacer(1, 6)]
        bullets = section.get("bullets") or []
        if bullets:
            story.append(
                ListFlowable([ListItem(Paragraph(escape(b), styles["BodyText"])) for b in bullets], bulletType="bullet")
            )
        table = section.get("table")
        if table:
            data = ([[_cell(h) for h in table.get("headers") or []]] if table.get("headers") else []) + [
                [_cell(c) for c in row] for row in table.get("rows") or []
            ]
            if data:
                grid = Table(data)
                grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)]))
                story += [grid, Spacer(1, 6)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=title or "").build(story)
    return buf.getvalue()


def _col(index: int) -> str:
    name = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        name = chr(65 + rem) + name
    return name


def render_xlsx(headers: list[str], rows: list[list[Scalar]]) -> bytes:
    if library_available("openpyxl"):
        import openpyxl  # type: ignore[import-untyped,unused-ignore]

        book = openpyxl.Workbook()
        sheet = book.active
        if headers:
            sheet.append(headers)
        for row in rows:
            sheet.append(list(row))
        buf = io.BytesIO()
        book.save(buf)
        return buf.getvalue()
    # Minimal SpreadsheetML (inline strings), readable by Excel, LibreOffice and openpyxl.
    lines = []
    table: list[list[Scalar]] = [list(headers)] if headers else []
    for r, values in enumerate(table + rows, start=1):
        cells = []
        for c, value in enumerate(values):
            ref = f"{_col(c)}{r}"
            if isinstance(value, bool) or value is None or not isinstance(value, int | float):
                text = escape(_cell(value))
                cells.append(f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{text}</t></is></c>')
            else:
                cells.append(f'<c r="{ref}"><v>{value}</v></c>')
        lines.append(f'<row r="{r}">{"".join(cells)}</row>')
    sheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        + "".join(lines)
        + "</sheetData></worksheet>"
    )
    files = {
        "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        "</Types>",
        "_rels/.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>",
        "xl/workbook.xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        "</Relationships>",
        "xl/worksheets/sheet1.xml": sheet,
    }
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, body in files.items():
            archive.writestr(name, body)
    return buf.getvalue()


def render_pptx(title: str | None, sections: list[dict[str, Any]]) -> bytes:
    import pptx  # type: ignore[import-not-found,unused-ignore]

    deck = pptx.Presentation()
    if title:
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = title
    for section in sections:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = section.get("heading") or ""
        body = slide.placeholders[1].text_frame
        items = list(section.get("paragraphs") or []) + list(section.get("bullets") or [])
        body.text = items[0] if items else ""
        for item in items[1:]:
            body.add_paragraph().text = item
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


_REQUIRES = {"docx": "docx", "pdf": "reportlab", "pptx": "pptx"}

_SECTION_SCHEMA: dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "properties": {
        "heading": {"type": "string"},
        "paragraphs": {"type": "array", "items": {"type": "string"}},
        "bullets": {"type": "array", "items": {"type": "string"}},
        "table": {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "headers": {"type": "array", "items": {"type": "string"}},
                "rows": {
                    "type": "array",
                    "items": {"type": "array", "items": {"type": ["string", "number", "boolean", "null"]}},
                },
            },
        },
    },
}


class DocumentsTool(ZTool):
    name = "documents"
    kind = "documents"
    description = (
        "Create a document as a job artifact. format: md, csv, json, mermaid, svg, docx, pdf, xlsx or pptx. "
        "Give `content` for md/json/mermaid/svg, `sections` (heading, paragraphs, bullets, table) for "
        "md/docx/pdf/pptx, and `headers` + `rows` for csv/xlsx. The file is saved as documents/<name>.<ext>."
    )
    input_schema: dict[str, Any] = {
        "type": "object",
        "additionalProperties": False,
        "required": ["format", "name"],
        "properties": {
            "format": {"type": "string", "enum": list(EXTENSIONS)},
            "name": {
                "type": "string",
                "pattern": r"^[A-Za-z0-9][A-Za-z0-9_-]{0,80}$",
                "description": "File name without extension.",
            },
            "title": {"type": "string"},
            "content": {"type": "string"},
            "sections": {"type": "array", "items": _SECTION_SCHEMA},
            "headers": {"type": "array", "items": {"type": "string"}},
            "rows": {
                "type": "array",
                "items": {"type": "array", "items": {"type": ["string", "number", "boolean", "null"]}},
            },
        },
    }

    def _render(self, fmt: str, args: dict[str, Any]) -> bytes:
        title = args.get("title")
        sections = _sections(args)
        content = args.get("content")
        if fmt in ("md", "json", "mermaid", "svg") and content is None and not (fmt == "md" and sections):
            raise ArgumentError([FieldError(path="content", message=f"'content' is required for format {fmt}")])
        if fmt in ("docx", "pdf", "pptx") and not sections and not title:
            raise ArgumentError(
                [FieldError(path="sections", message=f"'sections' or 'title' is required for format {fmt}")]
            )
        if fmt in ("csv", "xlsx") and not args.get("rows") and not args.get("headers"):
            raise ArgumentError([FieldError(path="rows", message=f"'headers' or 'rows' is required for format {fmt}")])
        if fmt == "md":
            return (content if content is not None else render_markdown(title, sections)).encode("utf-8")
        if fmt == "json":
            try:
                return json.dumps(json.loads(str(content)), indent=2, ensure_ascii=False).encode("utf-8")
            except json.JSONDecodeError as exc:
                raise ArgumentError([FieldError(path="content", message=f"not valid JSON: {exc}")]) from exc
        if fmt == "mermaid":
            return str(content).encode("utf-8")
        if fmt == "svg":
            text = str(content).strip()
            if not re.match(r"^(<\?xml[^>]*>\s*)?<svg\b", text) or _SVG_FORBIDDEN.search(text):
                raise ArgumentError(
                    [
                        FieldError(
                            path="content",
                            message="must be a static <svg> document (no scripts, event handlers or remote links)",
                        )
                    ]
                )
            return text.encode("utf-8")
        if fmt == "csv":
            return render_csv(list(args.get("headers") or []), list(args.get("rows") or []))
        if fmt == "xlsx":
            return render_xlsx(list(args.get("headers") or []), list(args.get("rows") or []))
        if fmt == "docx":
            return render_docx(title, sections)
        if fmt == "pdf":
            return render_pdf(title, sections)
        return render_pptx(title, sections)

    def run(self, args: dict[str, Any], ctx: ToolContext) -> ToolResult:
        fmt = str(args["format"])
        module = _REQUIRES.get(fmt)
        if module and not library_available(module):
            return ToolResult(
                text=f"Format {fmt} is not available in this runner image (missing library {module}). Use another format.",
                ok=False,
            )
        filename = f"{args['name']}.{EXTENSIONS[fmt]}"
        rel = f"documents/{filename}"
        try:
            target = ctx.workspace.resolve(rel, for_write=True)
        except JailError as exc:
            return ctx.jail_denial(rel, exc)
        data = self._render(fmt, args)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(data)
        ctx.progress.append(f"documents:{rel}")
        try:
            artifact_id = ctx.artifacts.upload(
                _KIND.get(fmt, ArtifactKind.document), filename, data, _CONTENT_TYPES.get(fmt)
            )
        except (ApiError, ArtifactRefused) as exc:
            return ToolResult(text=f"Saved {rel} ({len(data)} bytes) but the artifact upload failed: {exc}", ok=False)
        return ToolResult(
            text=f"Created {rel} ({len(data)} bytes), artifact {artifact_id}.",
            summary=f"created {rel}",
            artifact_ref=artifact_id,
        )
